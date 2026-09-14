"""Gera o enunciado da versão mínima de regressão logística.

Requer python-docx. Execute a partir de qualquer diretório:
    python scripts/create_task2_minimal_docx.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from create_task1_docx import (
    GRAY,
    LIGHT_TEAL,
    NAVY,
    TEAL,
    WHITE,
    add_bullet,
    add_callout,
    add_hyperlink,
    add_page_number,
    add_text,
    configure_section,
    set_cell_margins,
    set_cell_shading,
    set_repeat_table_header,
    style_document,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "tarefas" / "Tarefa_2_Aprendizado_Supervisionado_Versao_Minima.docx"
NOTEBOOK = "02_aprendizado_supervisionado_versão_minima.ipynb"
SOURCE = "https://archive.ics.uci.edu/dataset/891/cdc+diabetes+health+indicators"


def item(document, title, instructions):
    document.add_heading(title, level=2)
    for instruction in instructions:
        add_bullet(document, instruction)


def formula(document, text):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    return p


def comparison_table(document):
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [3.2, 8.4, 5.8]
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for cell, title in zip(table.rows[0].cells, ["Variável", "Contraste solicitado", "Cálculo da OR"]):
        set_cell_shading(cell, TEAL)
        add_text(cell.paragraphs[0], title, bold=True, color=WHITE, size=9.5)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("BMI (IMC)", "+1 kg/m²", "exp(b_BMI)"),
        ("BMI (IMC)", "+5 kg/m²", "exp(5 × b_BMI)"),
        ("HighBP", "1 (com pressão alta) versus 0 (sem)", "exp(b_HighBP)"),
    ]
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            set_cell_shading(cell, LIGHT_TEAL)
            add_text(cell.paragraphs[0], value, size=9.5)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            set_cell_margins(cell, top=95, bottom=95, start=120, end=120)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    return table


def build_document():
    document = Document()
    style_document(document)
    section = document.sections[0]
    configure_section(section)
    document.styles["Title"].font.size = Pt(27)
    document.styles["Subtitle"].font.size = Pt(12)
    document.styles["Subtitle"].paragraph_format.space_after = Pt(11)
    document.styles["Normal"].paragraph_format.keep_together = True
    document.styles["List Bullet"].paragraph_format.keep_together = True
    for style_name in ("Normal", "List Bullet"):
        props = document.styles[style_name].element.get_or_add_rPr()
        language = OxmlElement("w:lang")
        language.set(qn("w:val"), "pt-BR")
        props.append(language)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "APRENDIZADO DE MÁQUINA PARA SAÚDE", bold=True, color=TEAL, size=8)
    footer = section.footer.add_table(rows=1, cols=2, width=Cm(17.4))
    footer.columns[0].width = Cm(13.4)
    footer.columns[1].width = Cm(4)
    add_text(footer.cell(0, 0).paragraphs[0], "Tarefa 2 • Regressão logística • Versão mínima", color=GRAY, size=8)
    add_page_number(footer.cell(0, 1).paragraphs[0])

    p = document.add_paragraph()
    add_text(p, "TAREFA 2  |  VERSÃO MÍNIMA", bold=True, color=TEAL, size=10)
    document.add_heading("Regressão logística em saúde", level=0)
    document.add_paragraph("Desempenho, odds e odds ratio", style="Subtitle")
    add_callout(
        document,
        "Proposta da atividade",
        "Você utilizará indicadores de saúde do CDC Diabetes Health Indicators (UCI 891) "
        "para ajustar uma regressão logística, avaliar seu desempenho e interpretar seus "
        "coeficientes. Ao final, deverá distinguir probabilidade, odds e odds ratio e "
        "explicar os contrastes estimados pelo modelo. Material com finalidade exclusivamente educacional.",
    )
    document.add_heading("Material e configurações", level=1)
    p = document.add_paragraph("Notebook-base: ")
    add_text(p, NOTEBOOK, bold=True, size=9)
    document.add_paragraph(
        "Use o arquivo fornecido com esta atividade. Abra-o no Google Colab ou no Jupyter. "
        "São necessários numpy, pandas, scikit-learn e statsmodels, além de internet para carregar os dados."
    )
    p = document.add_paragraph("Fonte e dicionário: ")
    add_hyperlink(p, "CDC Diabetes Health Indicators — UCI 891", SOURCE)
    add_bullet(document, "Use todos os registros carregados pelo notebook e todos os 21 atributos; "
               "ID identifica os registros e não entra no modelo.")
    add_bullet(document, "Alvo Diabetes_binary: 0 = sem diabetes; 1 = pré-diabetes/diabetes. "
               "A classe positiva reúne as duas condições.")
    add_bullet(document, "Mantenha Age, GenHlth, Education e Income com os códigos numéricos "
               "do notebook. São categorias ordinais; considere essa simplificação ao interpretar.")
    document.add_heading("Roteiro de execução", level=1)
    document.add_paragraph(
        "Organize as respostas no próprio notebook, com código, saída e interpretação em "
        "células Markdown. Execute os itens na ordem e escreva as explicações em linguagem própria."
    )
    item(document, "1. Dados e divisão treino/teste", [
        "Carregue os dados e separe X e y. Reserve 20% para teste e 80% para treino, "
        "com stratify=y e random_state=42, como no notebook.",
        "Informe o número de registros em cada partição e a proporção da classe 1. "
        "Explique a finalidade da estratificação e do conjunto de teste.",
    ])
    item(document, "2. Padronização e treinamento", [
        "Ajuste make_pipeline(StandardScaler(), LogisticRegression(max_iter=1000)) "
        "somente no treino. Preserve os demais parâmetros padrão do modelo.",
        "Obtenha as classes previstas e as probabilidades da classe 1 no teste, com limiar "
        "padrão 0,5. Explique por que a média e a escala são aprendidas apenas no treino.",
    ])

    document.add_page_break()
    document.add_heading("Desempenho e odds", level=1)
    item(document, "3. Avaliação no teste", [
        "Apresente classification_report com três casas decimais e a ROC-AUC calculada "
        "com as probabilidades da classe 1.",
        "Identifique sensibilidade (recall da classe 1), especificidade (recall da classe 0), "
        "precisão e F1 da classe 1 e acurácia. Explique o que sensibilidade e especificidade medem.",
        "Comente o desempenho considerando o desbalanceamento das classes. Por que a "
        "acurácia, isoladamente, pode esconder erros relevantes?",
    ])
    item(document, "4. Coeficientes estimados", [
        "Apresente os coeficientes de todos os atributos na escala padronizada. Escolha "
        "um positivo e um negativo e explique a direção da associação com a log-odds da classe 1.",
        "Explique o significado de uma unidade padronizada. Por que a mudança de 0 para 1 "
        "em uma variável binária não equivale a um desvio-padrão?",
    ])
    item(document, "5. Das odds à probabilidade", [
        "Use os coeficientes e o intercepto já estimados, sem reajustar o modelo. Recupere "
        "mean_ e scale_ do StandardScaler e converta os parâmetros para a escala original.",
    ])
    document.add_paragraph(
        "Notação: βⱼ é o coeficiente padronizado; μⱼ e sⱼ são a média e a escala de treino; "
        "bⱼ e b₀ são os coeficientes e o intercepto convertidos. As somas percorrem todos os atributos."
    )
    formula(document, "bⱼ = βⱼ / sⱼ       b₀ = β₀ − Σⱼ bⱼ μⱼ")
    formula(document, "η = b₀ + Σⱼ bⱼ xⱼ       odds = exp(η)       p = odds / (1 + odds)")
    add_bullet(document, "Para os cinco primeiros registros de X_test, apresente ID, log-odds, "
               "odds, probabilidade calculada e probabilidade de predict_proba. Verifique se as "
               "duas probabilidades coincidem, admitindo pequenas diferenças numéricas.")
    add_bullet(document, "Escolha um registro e interprete suas odds e sua probabilidade. "
               "Explique por que o intercepto também precisa ser convertido e por que odds não é uma porcentagem.")
    add_callout(
        document,
        "Cuidados com o cálculo",
        "Use os parâmetros com sua precisão completa; arredonde somente as tabelas finais "
        "para quatro casas decimais. O teste fornece registros para a demonstração, mas "
        "não participa do ajuste dos parâmetros."
    )

    document.add_page_break()
    document.add_heading("Odds ratio, significância e entrega", level=1)
    item(document, "6. Odds ratio a partir dos coeficientes", [
        "Apresente, para todos os atributos, coeficiente padronizado, escala de treino, "
        "coeficiente original, OR para +1 unidade padronizada e OR para +1 unidade original.",
    ])
    formula(document, "OR(Δxⱼ) = exp(bⱼ × Δxⱼ)       Variação das odds (%) = 100 × (OR − 1)")
    document.add_paragraph(
        "Calcule e interprete os três contrastes abaixo, mantendo os demais atributos constantes. "
        "Na tabela, b_BMI e b_HighBP são os coeficientes na escala original."
    )
    comparison_table(document)
    add_bullet(document, "Para cada contraste, informe OR, variação percentual das odds, "
               "direção da associação e categoria ou incremento comparado. Use a classe 1 como desfecho.")
    add_bullet(document, "Explique a relação entre a OR para +5 unidades de IMC e a OR para +1. "
               "Interprete também uma OR menor que 1 da sua tabela, indicando o contraste correto.")
    add_bullet(document, "Explique por que exp(β_HighBP), calculada diretamente do coeficiente "
               "padronizado, não representa a comparação de 1 versus 0. Para Age, esclareça "
               "por que +1 no código não significa +1 ano.")
    add_bullet(document, "Na seção 5.3, ajuste também um GLM binomial sem penalização, com "
               "intercepto, os mesmos atributos nas unidades originais e somente os dados de treino. "
               "Use cov_type='HC0' e use_t=False para a inferência robusta por Wald bilateral.")
    add_bullet(document, "Apresente uma tabela final desse ajuste com coeficientes, erros-padrão, "
               "OR, IC95%, valores-p e indicação de p<0,05. Interprete HvyAlcoholConsump usando "
               "OR, IC95% e p em conjunto. Use notação científica para valores-p pequenos.")
    item(document, "7. Probabilidade, associação e limites", [
        "Em um exemplo hipotético, parta de p₀=0,20 e OR=2. Calcule as odds iniciais, "
        "as odds após o contraste e a nova probabilidade. A probabilidade dobrou? "
        "Explique por que uma OR, sozinha, não fornece a probabilidade de uma pessoa.",
        "Discuta, em um parágrafo, o caráter observacional e o autorrelato dos dados, a "
        "regularização e a simplificação dos códigos ordinais. Diferencie associação e causalidade. "
        "Os resultados deste exercício não validam uso clínico.",
    ])
    document.add_paragraph(
        "Os valores-p e IC95% pertencem ao ajuste sem penalização e às OR dessa tabela final. "
        "Os testes são exploratórios, sem correção para múltiplas comparações, e não incorporam "
        "o desenho amostral do BRFSS. Significância não demonstra causalidade ou importância clínica; "
        "p≥0,05 não prova ausência de associação."
    )
    document.add_heading("O que entregar", level=2)
    document.add_paragraph(
        "Entregue um arquivo .ipynb com identificação do(s) autor(es), os sete itens respondidos, "
        "código executado, tabelas e interpretações. Registre as versões de Python, numpy, pandas, "
        "scikit-learn e statsmodels e cite a fonte dos dados. Antes de entregar, reinicie o ambiente e execute "
        "todas as células em sequência."
    )
    document.add_paragraph(
        "Critérios de avaliação: execução reproduzível; separação correta entre treino e teste; "
        "cálculos e contrastes corretos; interpretação de odds e probabilidades; clareza na discussão dos limites."
    )
    p = document.add_paragraph()
    add_text(p, "Apoio conceitual: ", color=GRAY, size=8.5)
    add_hyperlink(p, "Penn State — regressão logística", "https://online.stat.psu.edu/stat462/node/207/")
    add_text(p, " • ", color=GRAY, size=8.5)
    add_hyperlink(p, "StandardScaler", "https://scikit-learn.org/stable/modules/generated/sklearn.preprocessing.StandardScaler.html")
    add_text(p, " • ", color=GRAY, size=8.5)
    add_hyperlink(p, "LogisticRegression", "https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.LogisticRegression.html")

    document.core_properties.title = "Tarefa 2 — Aprendizado supervisionado: versão mínima"
    document.core_properties.subject = "Enunciado: regressão logística, desempenho, odds e odds ratio"
    document.core_properties.author = "Flavio Luiz Seixas"
    document.core_properties.keywords = "regressão logística, odds, odds ratio, aprendizado supervisionado"
    document.core_properties.language = "pt-BR"
    return document


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    build_document().save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
