"""Gera notebook e enunciado Word da atividade de regressão logística em grupos.

Execute: python scripts/create_logistic_group_activity.py
Requer python-docx. O enunciado tem sua fonte editável em tarefas/*.md.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from create_task1_docx import TEAL, add_page_number, add_text, configure_section, style_document
from create_task3_minimal import add_inline, cell


ROOT = Path(__file__).resolve().parents[1]
NOTEBOOK = "extra_regressao_logistica_coracao.ipynb"
TASK = ROOT / "tarefas" / "Atividade_Extra_Regressao_Logistica_Coracao.md"
COLAB = "https://colab.research.google.com/github/flavioluizseixas/aprendizado-de-maquina-para-saude/blob/main/notebooks/"


def build_notebook() -> dict:
    cells = notebook_cells()
    for index, entry in enumerate(cells):
        entry["id"] = f"logistica-coracao-{index:02d}"
    return {
        "cells": cells,
        "metadata": {
            "colab": {"name": NOTEBOOK, "provenance": []},
            "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
            "language_info": {"name": "python"},
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }


def notebook_cells() -> list[dict]:
    return [
        cell("markdown", f"""
            # Regressão logística e doença cardíaca — atividade em grupos

            [![Abrir no Colab](https://colab.research.google.com/assets/colab-badge.svg)]({COLAB}{NOTEBOOK})

            **Grupos de 3–4 alunos | 60–75 minutos | Seis células de código.** Todos executam o mesmo modelo e discutem desempenho, coeficientes, odds e odds ratios. Ao final, cada grupo aprofunda um dos focos A–D.

            Execute em ordem no Colab ou em Jupyter com `numpy`, `pandas` e `scikit-learn`. Se necessário, use `%pip install numpy pandas scikit-learn` em uma célula adicional. O carregamento requer internet; não é necessário clonar o repositório.

            ## Fonte e licença

            [Heart Disease — Cleveland, UCI](https://archive.ics.uci.edu/dataset/45/heart+disease), Janosi, Steinbrunn, Pfisterer e Detrano (1989), licença CC BY 4.0. Usaremos 303 registros e cinco atributos escolhidos previamente para facilitar a interpretação.

            **Desfecho:** `y=0` quando `num=0`; `y=1` quando `num>0`, indicando presença de doença registrada na base. Não se trata de prever uma doença futura.

            | Atributo | Significado | Contraste para a OR |
            |---|---|---|
            | `age` | Idade em anos | +10 anos |
            | `trestbps` | Pressão arterial de repouso na admissão (mmHg) | +10 mmHg |
            | `chol` | Colesterol sérico (mg/dL) | +20 mg/dL |
            | `thalach` | Frequência cardíaca máxima atingida no exame (batimentos/minuto) | +10 batimentos/minuto |
            | `exang` | Angina induzida por exercício: 0 = não; 1 = sim | 1 versus 0 |

            Material com finalidade exclusivamente educacional. Esta amostra pequena e histórica não representa toda a população; o modelo não foi validado para decisões clínicas.

            ## 1. Dados e divisão treino/teste

            A base tem ausências em outros atributos. Verificamos apenas os cinco usados e o desfecho e removemos eventuais registros incompletos nesse subconjunto. Reservamos 25% para teste, com estratificação e semente 42. O índice serve somente para localizar linhas.
        """),
        cell("code", """
            import numpy as np
            import pandas as pd
            from IPython.display import display
            from sklearn.model_selection import train_test_split
            from sklearn.pipeline import make_pipeline
            from sklearn.preprocessing import StandardScaler
            from sklearn.linear_model import LogisticRegression
            from sklearn.metrics import (
                confusion_matrix, accuracy_score, recall_score,
                precision_score, f1_score, roc_auc_score,
            )

            atributos = ["age", "trestbps", "chol", "thalach", "exang"]
            dados = pd.read_csv("https://archive.ics.uci.edu/static/public/45/data.csv", na_values="?")
            display(dados[atributos + ["num"]].isna().sum().to_frame("Ausências"))
            dados = dados.dropna(subset=atributos + ["num"])
            X = dados[atributos]
            y = (dados["num"] > 0).astype(int)
            X_train, X_test, y_train, y_test = train_test_split(
                X, y, test_size=0.25, stratify=y, random_state=42
            )
            display(pd.DataFrame({
                "n": [len(y_train), len(y_test)],
                "Com doença (n)": [y_train.sum(), y_test.sum()],
                "Com doença (%)": [100 * y_train.mean(), 100 * y_test.mean()],
            }, index=["Treino", "Teste"]).round(1))
        """),
        cell("markdown", """
            ## 2. Estimar o modelo

            O pipeline aprende médias, escalas e coeficientes **somente no treino**. A regressão usa regularização L2, com `C=1`; seus coeficientes são estimados por otimização da log-verossimilhança penalizada. Não calcularemos valores-p ou intervalos de confiança nesta atividade.

            Fixamos o limiar em 0,5. Mantenha atributos, divisão e parâmetros iguais entre grupos; o teste será usado para avaliação, sem orientar ajustes do modelo.
        """),
        cell("code", """
            modelo = make_pipeline(
                StandardScaler(), LogisticRegression(C=1.0, max_iter=1000)
            )
            modelo.fit(X_train, y_train)
            y_prob = modelo.predict_proba(X_test)[:, 1]
            y_pred = (y_prob >= 0.5).astype(int)
        """),
        cell("markdown", """
            ## 3. Avaliar na base de teste

            Na matriz, **linhas = classe real** e **colunas = classe prevista**. Para a classe positiva (com doença), identifique VP, VN, FP e FN.

            - Sensibilidade = VP/(VP+FN); especificidade = VN/(VN+FP).
            - Precisão = VP/(VP+FP); F1 combina precisão e sensibilidade; acurácia é a proporção total de acertos.
            - ROC-AUC usa as probabilidades para avaliar a ordenação entre classes, sem depender de um único limiar; não garante probabilidades bem calibradas.

            A referência simples sempre prevê a classe mais frequente **do treino** e é avaliada no mesmo teste.
        """),
        cell("code", """
            matriz = confusion_matrix(y_test, y_pred, labels=[0, 1])
            display(pd.DataFrame(
                matriz, index=["Real: sem doença", "Real: com doença"],
                columns=["Previsto: sem doença", "Previsto: com doença"]
            ))
            metricas = pd.Series({
                "Acurácia": accuracy_score(y_test, y_pred),
                "Sensibilidade": recall_score(y_test, y_pred),
                "Especificidade": recall_score(y_test, y_pred, pos_label=0),
                "Precisão": precision_score(y_test, y_pred, zero_division=0),
                "F1": f1_score(y_test, y_pred, zero_division=0),
                "ROC-AUC": roc_auc_score(y_test, y_prob),
            }, name="Resultado no teste")
            display(metricas.round(3).to_frame())
            classe_majoritaria = y_train.mode().iloc[0]
            acuracia_referencia = (y_test == classe_majoritaria).mean()
            print(f"Referência (sempre classe {classe_majoritaria}): acurácia = {acuracia_referencia:.3f}")
        """),
        cell("markdown", """
            ## 4. Coeficientes nas duas escalas

            O modelo usa `zⱼ=(xⱼ−μⱼ)/sⱼ`, com média μⱼ e escala sⱼ aprendidas no treino. Seu coeficiente βⱼ se refere a uma unidade padronizada. Para voltar às unidades originais:

            **bⱼ = βⱼ/sⱼ** e **b₀ = β₀ − Σⱼ bⱼμⱼ**.

            Coeficiente positivo aumenta a log-odds estimada da classe 1; negativo a reduz, mantendo os demais atributos constantes. Coeficientes em unidades diferentes não devem ser comparados diretamente como importância. Em `exang`, um desvio-padrão não equivale à mudança de 0 para 1.

            A conversão reaproveita o modelo treinado e mantém sua regularização. O intercepto corresponde a todos os atributos originais iguais a zero, um perfil sem interpretação clínica útil aqui.
        """),
        cell("code", """
            padronizador, regressao = modelo[0], modelo[-1]
            coef_pad = pd.Series(regressao.coef_[0], index=atributos)
            coef_original = coef_pad / padronizador.scale_
            intercepto_original = regressao.intercept_[0] - (coef_original * padronizador.mean_).sum()

            coeficientes = pd.DataFrame({
                "Coeficiente padronizado": coef_pad,
                "Coeficiente original": coef_original,
            })
            display(coeficientes.round(4))
            print(f"Intercepto na escala original: {intercepto_original:.4f}")
        """),
        cell("markdown", """
            ## 5. Dos coeficientes às odds e à probabilidade

            Para um registro, **η = b₀ + Σⱼ bⱼxⱼ**, **odds = exp(η)** e **p = odds/(1+odds)**. Equivalentemente, odds = p/(1−p).

            Por exemplo, p=0,20 corresponde a odds=0,25, uma razão de 1 para 4 entre as probabilidades de presença e ausência. Odds não são porcentagens.

            Reconstruímos as probabilidades dos cinco primeiros registros do teste e as comparamos com `predict_proba`, sem refazer o ajuste. Os parâmetros completos são usados no cálculo; o arredondamento serve apenas à apresentação.
        """),
        cell("code", """
            perfis = X_test.head(5)
            log_odds = intercepto_original + perfis @ coef_original
            odds = np.exp(log_odds)
            prob_calculada = odds / (1 + odds)
            prob_pipeline = modelo.predict_proba(perfis)[:, 1]

            tabela_odds = perfis.assign(
                observado=y_test.loc[perfis.index], log_odds=log_odds,
                odds=odds, prob_calculada=prob_calculada, prob_pipeline=prob_pipeline,
            )
            display(tabela_odds.round(4))
            print("Probabilidades coincidem:", np.allclose(prob_calculada, prob_pipeline))
        """),
        cell("markdown", """
            ## 6. Odds ratios com contrastes definidos

            Mantendo os demais atributos constantes, **OR(Δxⱼ) = exp(bⱼ × Δxⱼ)**. A tabela usa os coeficientes na escala original e incrementos fáceis de discutir.

            **OR > 1:** odds maiores; **OR < 1:** odds menores. A variação percentual das odds é **100 × (OR−1)**. OR=1,50 indica odds 50% maiores, sem significar probabilidade 50% maior. OR próxima de 1 não estabelece ausência de associação estatística; não fizemos um teste de significância.

            Para `exang`, o contraste é presença versus ausência de angina induzida por exercício. Estas são associações condicionais do modelo; não demonstram que modificar um atributo causaria mudança na doença.
        """),
        cell("code", """
            contrastes = pd.DataFrame({
                "Atributo": atributos,
                "Comparação": ["+10 anos", "+10 mmHg", "+20 mg/dL", "+10 batimentos/min", "1 (sim) versus 0 (não)"],
                "Incremento": [10, 10, 20, 10, 1],
            })
            contrastes["OR"] = np.exp(
                contrastes["Atributo"].map(coef_original) * contrastes["Incremento"]
            )
            contrastes["Variação das odds (%)"] = 100 * (contrastes["OR"] - 1)
            display(contrastes.round(4))
        """),
        cell("markdown", """
            ## Discussão em grupos

            Todos respondem à parte comum. O professor distribui um foco A–D por grupo; focos podem ser repetidos.

            - **A — Erros:** explique FP e FN e discuta o efeito esperado de reduzir o limiar na sensibilidade e na especificidade, sem escolher um limiar pelo teste.
            - **B — Idade e pressão:** interprete as OR para +10 anos e +10 mmHg. Calcule `np.exp(coef_original["age"] * 20)` e compare com a OR de +10 anos elevada ao quadrado.
            - **C — Colesterol e frequência cardíaca:** interprete as OR para +20 mg/dL e +10 batimentos/minuto. Um coeficiente negativo demonstra proteção causal? Considere correlação entre atributos e seleção da amostra.
            - **D — Angina e probabilidade:** interprete a OR de `exang`. Para um perfil hipotético com `exang=0` e `p0 = 0.20`, calcule `or_exang = np.exp(coef_original["exang"])`, `odds1 = or_exang * p0 / (1 - p0)` e `p1 = odds1 / (1 + odds1)`, mantendo os demais atributos constantes. Compare `p1` com `or_exang * p0`: por que não são iguais? Não se trata do efeito de uma intervenção.

            ## Respostas do grupo

            **Participantes:** preencher. **Foco atribuído:** preencher.

            1. **Dados e ajuste:** registre tamanhos e casos positivos do treino/teste. Explique estratificação e padronização apenas no treino.
            2. **Desempenho:** registre VP, VN, FP, FN e as seis métricas. Confira sensibilidade e especificidade pela matriz e compare a acurácia com a referência.
            3. **Coeficientes:** interprete um positivo e um negativo, se houver. Diferencie as escalas e explique a conversão do intercepto.
            4. **Odds:** escolha um registro da tabela, informe p e odds e confira a conversão entre ambos.
            5. **OR:** interprete um contraste contínuo e o de `exang`, mantendo os demais atributos constantes.
            6. **Foco e limites:** responda à questão atribuída e descreva duas limitações do estudo.

            **Entrega:** um `.ipynb` por grupo com saídas e respostas em Markdown. Adicione uma célula para os cálculos do foco, se necessário. Prepare três minutos de apresentação com uma métrica, uma OR e uma limitação. Reinicie o ambiente e execute tudo antes de entregar.

            **Referências:** [regressão logística](https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.LogisticRegression.html) e [StandardScaler](https://scikit-learn.org/stable/modules/generated/sklearn.preprocessing.StandardScaler.html), scikit-learn.
        """),
    ]


def build_document() -> Document:
    document = Document()
    style_document(document)
    configure_section(document.sections[0])
    document.styles["Title"].font.size = Pt(23)
    document.styles["Heading 1"].font.size = Pt(14)
    document.styles["Heading 2"].font.size = Pt(11)
    for name in ("Normal", "List Bullet"):
        document.styles[name].paragraph_format.keep_together = True
    header = document.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "APRENDIZADO DE MÁQUINA PARA SAÚDE", bold=True, color=TEAL, size=8)
    add_page_number(document.sections[0].footer.paragraphs[0])
    document.core_properties.title = "Atividade em grupos — Regressão logística e doença cardíaca"
    document.core_properties.language = "pt-BR"
    for block in TASK.read_text(encoding="utf-8").strip().split("\n\n"):
        if block.startswith("#"):
            prefix, title = block.split(" ", 1)
            document.add_heading(title, level=len(prefix) - 1)
        elif block.startswith("- "):
            for line in block.splitlines():
                add_inline(document.add_paragraph(style="List Bullet"), line[2:])
        else:
            add_inline(document.add_paragraph(), block)
    return document


def main() -> None:
    notebook_path = ROOT / "notebooks" / NOTEBOOK
    notebook_path.write_text(
        json.dumps(build_notebook(), ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )
    build_document().save(TASK.with_suffix(".docx"))
    print(notebook_path)
    print(TASK.with_suffix(".docx"))


if __name__ == "__main__":
    main()
